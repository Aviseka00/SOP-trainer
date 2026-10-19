"""Local-first SOP storyboard and video composition engine. Python 3.11+."""
import io, json, re, shutil, subprocess, wave, asyncio, uuid, textwrap
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parent
for folder in ('approved_media','outputs','jobs'): (ROOT/folder).mkdir(exist_ok=True)
CATEGORIES={
 'Hand hygiene':['handwash','hand wash','soap','hand hygiene','rinse','scrub'],
 'Gowning':['gown','coverall','shoe cover','donning','doffing','ppe','mask'],
 'Disinfection':['sanitiz','disinfect','alcohol','wipe','cleaning'],
 'Equipment':['equipment','dryer','fbd','machine','calibrat','operate'],
 'Documentation':['logbook','record','document','sign','batch record'],
 'Safety':['hazard','spill','emergency','injury','alarm'],
 'Material handling':['dispens','weigh','transfer','container','material'],
 'Aseptic technique':['aseptic','sterile','laminar','biosafety','grade a']}
LAYOUTS=['Demonstration full screen','Presenter picture-in-picture','Presenter side-by-side','Presenter full screen']

def run(args):
 p=subprocess.run([str(a) for a in args],capture_output=True,text=True)
 if p.returncode: raise RuntimeError((p.stderr or p.stdout)[-2500:])
 return p.stdout

def ffmpeg():
 try:
  import imageio_ffmpeg
  return shutil.which('ffmpeg') or imageio_ffmpeg.get_ffmpeg_exe()
 except ImportError: return shutil.which('ffmpeg')

def duration(path):
 q=shutil.which('ffprobe')
 if q:
  try:return float(run([q,'-v','error','-show_entries','format=duration','-of','default=noprint_wrappers=1:nokey=1',path]).strip())
  except Exception:pass
 if str(path).lower().endswith('.wav'):
  with wave.open(str(path),'rb') as w:return w.getnframes()/w.getframerate()
 return None

def extract(filename,content):
 name=filename.lower()
 if name.endswith('.pdf'):
  import fitz
  with fitz.open(stream=content,filetype='pdf') as doc:return '\n'.join(p.get_text() for p in doc)
 if name.endswith('.docx'):
  from docx import Document
  doc=Document(io.BytesIO(content))
  return '\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells])
 return content.decode('utf-8',errors='replace')

def category(text):
 s=text.lower(); scores={k:sum(s.count(w) for w in v) for k,v in CATEGORIES.items()}
 return max(scores,key=scores.get) if max(scores.values(),default=0)>0 else 'General'

def storyboard(text):
 chunks=[]
 for para in re.split(r'\n+',text):
  para=para.strip(' \t•-')
  if not para:continue
  for part in re.split(r'(?<=[.!?])\s+(?=[A-Z])',para):
   if part.strip():chunks.append(part.strip())
 return [{'narration':c,'visual':category(c),'layout':'Presenter picture-in-picture' if i%3!=1 else 'Demonstration full screen','reference':f'Source segment {i+1}','approved':False,'media':'','seconds':0} for i,c in enumerate(chunks[:80])]

def font(sz):
 for p in ['C:/Windows/Fonts/arial.ttf','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf']:
  if Path(p).exists():return ImageFont.truetype(p,sz)
 return ImageFont.load_default()

def wrap(draw,text,ft,max_width):
 words=text.split();lines=[];cur=''
 for w in words:
  candidate=(cur+' '+w).strip()
  if draw.textbbox((0,0),candidate,font=ft)[2]>max_width and cur:lines.append(cur);cur=w
  else:cur=candidate
 if cur:lines.append(cur)
 return lines

def placeholder(row,dest,title):
 im=Image.new('RGB',(1280,720),(14,26,44));d=ImageDraw.Draw(im)
 d.rounded_rectangle((28,28,1252,692),radius=22,outline=(47,115,142),width=3)
 d.text((64,58),title[:55],font=font(32),fill=(124,224,230))
 d.text((64,145),row['visual'],font=font(48),fill='white')
 for j,line in enumerate(wrap(d,row['narration'],font(28),1100)[:5]):d.text((65,270+j*42),line,font=font(28),fill=(217,228,236))
 d.text((65,627),'VISUAL PLACEHOLDER — NOT A PROCEDURAL DEMONSTRATION',font=font(21),fill=(255,195,121))
 im.save(dest)

async def edge_tts(text,path,voice,rate):
 import edge_tts
 await edge_tts.Communicate(text,voice,rate=rate).save(str(path))

def synth(text,path,engine,voice,rate):
 if engine=='Edge (online)':asyncio.run(edge_tts(text,path,voice,rate))
 else:
  import pyttsx3
  t=pyttsx3.init();t.setProperty('rate',int(rate));t.save_to_file(text,str(path));t.runAndWait();t.stop()
 if not path.exists() or path.stat().st_size<200:raise RuntimeError('Speech synthesis failed; check installed Windows voices or use Edge online.')

def stamp(t):
 ms=round(t*1000);return f'{ms//3600000:02}:{ms//60000%60:02}:{ms//1000%60:02},{ms%1000:03}'

def render(rows,title,engine,voice,rate,presenter_path=None,media_map=None,output_dir=None,progress=None,preexisting_audio=None):
 """Render verified scenes with voice, contextual assets, presenter video (not synthetic lip sync)."""
 ff=ffmpeg()
 if not ff:raise RuntimeError('FFmpeg is missing.')
 job=ROOT/'jobs'/uuid.uuid4().hex[:10];job.mkdir(parents=True)
 out=Path(output_dir) if output_dir else ROOT/'outputs';out.mkdir(parents=True,exist_ok=True)
 media_map=media_map or {};clips=[];subs=[];timing=[];offset=0.
 for i,row in enumerate(rows):
  if not row.get('approved'):raise ValueError(f'Scene {i+1} needs QA review.')
  audio=Path(preexisting_audio[i]) if preexisting_audio else job/f'audio_{i:03}.{ "mp3" if engine=="Edge (online)" else "wav"}'
  if not preexisting_audio:synth(row['narration'],audio,engine,voice,rate)
  d=duration(audio)
  if d is None:
   probe=job/f'probe_{i:03}.wav';run([ff,'-y','-i',audio,'-ac','1','-ar','16000',probe]);d=duration(probe)
  if not d:raise RuntimeError(f'Unable to read speech duration for scene {i+1}')
  d=max(0.6,d+0.15)
  media=row.get('media') or media_map.get(row['visual'],'')
  src=Path(media) if media and Path(media).exists() else None
  if not src:
   src=job/f'placeholder_{i:03}.png';placeholder(row,src,title)
  video=job/f'clip_{i:03}.mp4'
  is_vid=src.suffix.lower() in ('.mp4','.mov','.mkv','.webm','.avi')
  # Fixed 1280x720 canvas; any real presenter is an existing recorded clip, not generated facial animation.
  cmd=[ff,'-y','-loglevel','error']
  cmd+=['-stream_loop','-1','-i',src] if is_vid else ['-loop','1','-framerate','25','-i',src]
  use_presenter=bool(presenter_path and Path(presenter_path).exists() and row['layout']!='Demonstration full screen')
  if use_presenter:cmd+=['-stream_loop','-1','-i',presenter_path]
  cmd+=['-i',audio]
  base='[0:v]scale=1280:720:force_original_aspect_ratio=increase,crop=1280:720,setsar=1,fps=25,format=yuv420p[bg]'
  if use_presenter:
   if row['layout']=='Presenter full screen':
    flt='[1:v]scale=1280:720:force_original_aspect_ratio=increase,crop=1280:720,setsar=1,fps=25,format=yuv420p[v]'
   elif row['layout']=='Presenter side-by-side':
    flt=base+';[1:v]scale=560:720:force_original_aspect_ratio=increase,crop=560:720,setsar=1,fps=25[p];[bg]crop=720:720:560:0[dem];[p][dem]hstack=inputs=2,format=yuv420p[v]'
   else:
    flt=base+';[1:v]scale=390:220:force_original_aspect_ratio=increase,crop=390:220,setsar=1,fps=25[p];[bg][p]overlay=W-w-28:H-h-28:shortest=0,format=yuv420p[v]'
  else:flt=base+';[bg]copy[v]'
  audio_index=2 if use_presenter else 1
  cmd+=['-filter_complex',flt,'-map','[v]','-map',f'{audio_index}:a:0','-t',f'{d:.3f}','-c:v','libx264','-preset','veryfast','-crf','24','-pix_fmt','yuv420p','-c:a','aac','-ar','44100','-b:a','128k','-shortest',video]
  run(cmd);clips.append(video)
  subs.append(f'{i+1}\n{stamp(offset)} --> {stamp(offset+d)}\n{row["narration"]}\n')
  timing.append({'scene':i+1,'start':round(offset,3),'end':round(offset+d,3),'visual':row['visual'],'layout':row['layout'],'reference':row['reference'],'media':str(src)})
  offset+=d
  if progress:progress((i+1)/len(rows))
 playlist=job/'concat.txt';playlist.write_text(''.join("file '"+p.name+"'\n" for p in clips),encoding='utf-8')
 mp4=out/f'SOP_Training_{job.name}.mp4';run([ff,'-y','-loglevel','error','-f','concat','-safe','0','-i',playlist,'-c','copy',mp4])
 srt=out/f'SOP_Training_{job.name}.srt';srt.write_text('\n'.join(subs),encoding='utf-8')
 manifest=out/f'SOP_Training_{job.name}.json';manifest.write_text(json.dumps({'title':title,'scenes':rows,'timeline':timing,'warning':'Media and narration require independent GMP approval. Presenter footage is not automatically lip-synced.'},indent=2,ensure_ascii=False),encoding='utf-8')
 return mp4,srt,manifest
