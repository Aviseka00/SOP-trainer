import asyncio, io, json, re, subprocess, shutil, wave, uuid
from pathlib import Path
import streamlit as st
import fitz
from docx import Document
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parent
MEDIA=ROOT/'approved_media'; MEDIA.mkdir(exist_ok=True)
OUT=ROOT/'outputs'; OUT.mkdir(exist_ok=True)
CATS={'Hand hygiene':['soap','wash hands','handwashing','hand hygiene','rinse'], 'Gowning':['gown','coverall','shoe cover','ppe','mask'], 'Sanitization':['sanitiz','disinfect','alcohol'], 'Equipment':['equipment','dryer','fbd','machine','start'], 'Documentation':['logbook','record','document','sign'], 'Safety':['hazard','spill','emergency','injury'], 'Material handling':['dispens','weigh','transfer','container']}

def run(*args):
    p=subprocess.run(list(map(str,args)),capture_output=True,text=True)
    if p.returncode: raise RuntimeError((p.stderr or p.stdout)[-1800:])
    return p.stdout

def ffmpeg():
    try:
        import imageio_ffmpeg
        return shutil.which('ffmpeg') or imageio_ffmpeg.get_ffmpeg_exe()
    except ImportError: return shutil.which('ffmpeg')

def probe_duration(p):
    q=shutil.which('ffprobe')
    if q:
        try: return float(run(q,'-v','error','-show_entries','format=duration','-of','default=noprint_wrappers=1:nokey=1',p).strip())
        except Exception: pass
    if str(p).endswith('.wav'):
        with wave.open(str(p),'rb') as w: return w.getnframes()/w.getframerate()
    return None

def extract(f):
    b=f.getvalue(); n=f.name.lower()
    if n.endswith('.pdf'):
        with fitz.open(stream=b,filetype='pdf') as doc: return '\n'.join(page.get_text() for page in doc)
    if n.endswith('.docx'):
        d=Document(io.BytesIO(b)); return '\n'.join([p.text for p in d.paragraphs]+[c.text for t in d.tables for r in t.rows for c in r.cells])
    return b.decode('utf-8',errors='replace')

def category(t):
    x=t.lower(); score={k:sum(x.count(w) for w in v) for k,v in CATS.items()}
    return max(score,key=score.get) if max(score.values(),default=0) else 'General'

def scenes(t):
    parts=[p.strip(' \t-•') for p in re.split(r'\n+|(?<=[.!?])\s+(?=[A-Z])',t) if len(p.strip())>15]
    return [{'narration':p,'visual':category(p),'approved':False,'source':str(i+1)} for i,p in enumerate(parts[:40])]

def font(sz):
    for p in ['C:/Windows/Fonts/arial.ttf','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf']:
        if Path(p).exists(): return ImageFont.truetype(p,sz)
    return ImageFont.load_default()

def card(scene,path,heading):
    im=Image.new('RGB',(1280,720),(12,25,42)); d=ImageDraw.Draw(im)
    d.rounded_rectangle((35,35,1245,685),radius=25,outline=(66,130,153),width=3)
    d.text((75,75),heading[:65],font=font(30),fill=(111,220,230))
    d.text((75,155),scene['visual'],font=font(52),fill='white')
    # Only approved stills/clips demonstrate steps; fallback is explicitly non-demonstrative.
    d.text((75,280),'VISUAL NOT YET PROVIDED',font=font(32),fill=(250,193,112))
    d.text((75,342),'Attach a QA-approved demonstration clip or image.',font=font(26),fill='white')
    d.text((75,630),'TRAINING PREVIEW • NOT AN APPROVED PROCEDURAL DEMONSTRATION',font=font(20),fill=(180,199,210))
    im.save(path)

async def edge_speak(text,path,voice,rate):
    import edge_tts
    await edge_tts.Communicate(text,voice,rate=rate).save(str(path))

def speak(text,path,engine,voice,rate):
    if engine=='Edge TTS (internet required)':
        asyncio.run(edge_speak(text,path,voice,rate)); return
    import pyttsx3
    e=pyttsx3.init(); e.setProperty('rate',155)
    e.save_to_file(text,str(path)); e.runAndWait(); e.stop()
    if not path.exists() or path.stat().st_size<100: raise RuntimeError('Windows voice failed. Try Edge TTS.')

def srt_stamp(s):
    ms=round(s*1000); return f'{ms//3600000:02}:{(ms//60000)%60:02}:{(ms//1000)%60:02},{ms%1000:03}'

def render(rows,media_map,engine,voice,rate,title,job):
    f=ffmpeg()
    if not f: raise RuntimeError('FFmpeg missing. Install ffmpeg or imageio-ffmpeg.')
    job.mkdir(exist_ok=True); concat=[]; subtitles=[]; offset=0
    for i,row in enumerate(rows):
        audio=job/f'audio_{i:03}.mp3' if engine.startswith('Edge') else job/f'audio_{i:03}.wav'
        speak(row['narration'],audio,engine,voice,rate)
        dur=probe_duration(audio)
        if dur is None:
            # ffmpeg's own progress duration probe via decoding
            wav=job/f'probe_{i:03}.wav'
            run(f,'-y','-i',audio,'-ar','16000','-ac','1',wav)
            dur=probe_duration(wav)
        dur=max(1.0,dur+0.2)
        asset=media_map.get(row['visual'])
        if asset and Path(asset).exists(): source=Path(asset)
        else:
            source=job/f'card_{i:03}.png'; card(row,source,title)
        clip=job/f'clip_{i:03}.mp4'
        filt='scale=1280:720:force_original_aspect_ratio=decrease,pad=1280:720:(ow-iw)/2:(oh-ih)/2,setsar=1,fps=25,format=yuv420p'
        if source.suffix.lower() in ['.mp4','.mov','.mkv','.webm']:
            run(f,'-y','-stream_loop','-1','-i',source,'-i',audio,'-t',f'{dur:.3f}','-vf',filt,'-map','0:v:0','-map','1:a:0','-c:v','libx264','-preset','veryfast','-c:a','aac','-ar','44100','-shortest',clip)
        else:
            run(f,'-y','-loop','1','-framerate','25','-i',source,'-i',audio,'-t',f'{dur:.3f}','-vf',filt,'-map','0:v:0','-map','1:a:0','-c:v','libx264','-preset','veryfast','-c:a','aac','-ar','44100','-shortest',clip)
        concat.append(clip)
        subtitles.append(f'{i+1}\n{srt_stamp(offset)} --> {srt_stamp(offset+dur)}\n{row["narration"]}\n')
        offset+=dur
    playlist=job/'playlist.txt'
    playlist.write_text(''.join("file '"+p.name+"'\n" for p in concat),encoding='utf-8')
    dest=OUT/f'{job.name}.mp4'
    run(f,'-y','-f','concat','-safe','0','-i',playlist,'-c','copy',dest)
    srt=OUT/f'{job.name}.srt'; srt.write_text('\n'.join(subtitles),encoding='utf-8')
    return dest,srt

st.set_page_config(page_title='SOP Trainer Studio Pro',layout='wide',page_icon='🎬')
st.title('🎬 SOP Trainer Studio Pro')
st.caption('SOP → reviewable scenes → multilingual voice → synchronized approved visuals → MP4 | Lip-synced avatar is an optional external stage, not simulated here.')
with st.sidebar:
    st.header('Narration')
    engine=st.selectbox('Voice engine',['Edge TTS (internet required)','Windows offline TTS'])
    voices={'English (India) • Female':'en-IN-NeerjaNeural','English (India) • Male':'en-IN-PrabhatNeural','Hindi • Female':'hi-IN-SwaraNeural','Hindi • Male':'hi-IN-MadhurNeural','Odia • Female':'or-IN-SubhasiniNeural','Odia • Male':'or-IN-SukantNeural','English (US) • Female':'en-US-JennyNeural'}
    language=st.selectbox('Voice / language',list(voices))
    voice=voices[language]; rate=st.select_slider('Speaking pace',options=['-20%','-10%','+0%','+10%','+20%'],value='+0%')
    st.info('Edge TTS uses an online Microsoft speech service. Offline Windows TTS depends on voices installed in Windows; it may not support Hindi/Odia.')
    st.warning('Text is spoken as entered. Selecting another language voice does NOT translate the SOP.')

title=st.text_input('SOP title','Personnel Hygiene')
file=st.file_uploader('Upload SOP (PDF, DOCX, TXT)',type=['pdf','docx','txt'])
source=extract(file) if file else ''
source=st.text_area('Or paste / edit SOP text',value=source,height=190)
if st.button('Build editable storyboard',type='primary'):
    st.session_state['rows']=scenes(source)
    st.session_state['media']={}
if 'rows' in st.session_state:
    st.subheader('Storyboard & GMP approval')
    st.warning('Review exact procedure, sequence, narration and visuals against the effective approved SOP. Generated text and media are not GMP validated.')
    rows=st.session_state['rows']; media=st.session_state['media']
    for i,r in enumerate(rows):
        with st.expander(f'Scene {i+1}: {r["visual"]}',expanded=i==0):
            r['narration']=st.text_area('Narration (must match approved SOP)',r['narration'],key=f'n_{i}')
            r['visual']=st.selectbox('Visual category',list(CATS)+['General'],index=(list(CATS)+['General']).index(r['visual']) if r['visual'] in list(CATS)+['General'] else -1,key=f'v_{i}')
            r['source']=st.text_input('SOP section / reference',r['source'],key=f's_{i}')
            r['approved']=st.checkbox('I reviewed this scene against the approved SOP',r['approved'],key=f'a_{i}')
    st.subheader('Approved demonstration library')
    st.caption('Attach real footage, animation, or still image for each category. The app automatically switches visuals when narration moves to that category.')
    for cat in sorted(set(r['visual'] for r in rows)):
        upload=st.file_uploader(f'{cat}: approved image or video',type=['png','jpg','jpeg','mp4','mov','webm'],key=f'upload_{cat}')
        if upload:
            dest=MEDIA/(re.sub(r'[^a-z0-9]+','_',cat.lower())+'_'+Path(upload.name).name)
            dest.write_bytes(upload.getvalue()); media[cat]=str(dest)
        if cat in media: st.caption('Selected: '+Path(media[cat]).name)
    st.download_button('Export scene plan JSON',json.dumps(rows,indent=2,ensure_ascii=False),file_name='sop_storyboard.json')
    if st.button('Render reviewed training video',type='primary',disabled=not rows or not all(r['approved'] for r in rows)):
        try:
            with st.spinner('Generating speech and rendering scene-synchronized video...'):
                job=ROOT/'jobs'/str(uuid.uuid4())[:8]
                video,srt=render(rows,media,engine,voice,rate,title,job)
            st.success('Rendered. This output does not include a photorealistic lip-synced face.')
            st.video(str(video))
            st.download_button('Download MP4',video.read_bytes(),file_name=video.name,mime='video/mp4')
            st.download_button('Download SRT subtitles',srt.read_bytes(),file_name=srt.name)
        except Exception as e: st.error(str(e))
