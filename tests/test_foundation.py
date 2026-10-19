import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from sop_studio.ingestion import storyboard, metadata, extract, contradiction_flags
from sop_studio.schema import approve, validate_project, validate_scene, file_hash
from sop_studio.storage import ProjectStore
from sop_studio.media import MediaLibrary
from sop_studio.composition import preflight, render
from sop_studio.process import run, RenderCancelled
from sop_studio.speech import pronunciation

TEST_TEMP = Path(__file__).resolve().parent.parent / '.test_tmp'
TEST_TEMP.mkdir(exist_ok=True)
tempfile.tempdir = str(TEST_TEMP)

class FoundationTests(unittest.TestCase):
    def test_preserves_more_than_eighty_lines(self):
        text = '\n'.join(f'3.{i+1} Apply 2.5 mL of approved soap.' for i in range(100))
        rows = storyboard(text)
        self.assertEqual(100, len(rows))
        self.assertEqual(text.splitlines()[89], rows[89]['original_instruction'])
        self.assertIn('3.90', rows[89]['reference'])

    def test_metadata_and_prohibition(self):
        text = 'Title: Test\nSOP ID: ABC-001\nRevision: 03\nEffective date: 2026-01-01\n2.1 Do not use equipment EQ-12.'
        self.assertEqual('03', metadata(text)['revision'])
        self.assertEqual('ABC-001', metadata(text)['sop_number'])
        row = storyboard(text)[-1]
        self.assertEqual('prohibited', row['safety_classification'])
        self.assertEqual(['EQ-12'], row['entities']['equipment'])
        self.assertTrue(contradiction_flags(row['narration'], 'Use equipment.'))

    def test_empty_and_invalid_documents(self):
        for filename, content in [('a.txt', b''), ('a.txt', b'\xffbad'), ('a.exe', b'data'), ('a.pdf', b'invalid')]:
            with self.assertRaises(Exception):
                extract(filename, content)

    def test_docx_keeps_table_order(self):
        from docx import Document
        document = Document()
        document.add_paragraph('Step one')
        document.add_table(rows=1, cols=1).cell(0, 0).text = 'Step two'
        document.add_paragraph('Step three')
        stream = io.BytesIO()
        document.save(stream)
        self.assertEqual('Step one\nStep two\nStep three', extract('sop.docx', stream.getvalue()))

    def test_pdf_and_scanned_rejection(self):
        import fitz
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((30, 30), '2.1 Apply approved soap.')
        self.assertIn('2.1 Apply approved soap.', extract('sop.pdf', doc.tobytes()))
        blank = fitz.open()
        blank.new_page()
        with self.assertRaisesRegex(ValueError, 'OCR'):
            extract('scan.pdf', blank.tobytes())

    def test_approval_is_bound_to_content(self):
        row = storyboard('Do not use damaged equipment.')[0]
        approve(row, 'QA reviewer')
        self.assertTrue(validate_scene(row)['approved'])
        row['narration'] = 'Use damaged equipment.'
        self.assertFalse(validate_scene(row)['approved'])
        with self.assertRaises(ValueError):
            preflight([row], {}, None, None)

    def test_approval_requires_reviewer(self):
        with self.assertRaises(ValueError):
            approve(storyboard('Wash hands.')[0], '')

    def test_validation_bad_types(self):
        for scene in [{'narration': ''}, {'narration': 'x', 'media_offset': float('nan')}, {'narration': 'x', 'layout': 'unknown'}]:
            with self.assertRaises(ValueError):
                validate_scene(scene)
        with self.assertRaises(ValueError):
            validate_project({'scenes': 'invalid'})

    def test_atomic_store_and_history(self):
        with tempfile.TemporaryDirectory() as folder:
            store = ProjectStore(Path(folder))
            p = store.save({'scenes': storyboard('Wash hands.'), 'title': 'First'})
            p['title'] = 'Second'
            p = store.save(p)
            self.assertEqual('Second', store.load(p['project_id'])['title'])
            self.assertEqual(2, len(p['history']))
            self.assertEqual(1, len(list((Path(folder) / p['project_id'] / 'history').glob('*.json'))))
            with self.assertRaises(ValueError):
                store.load('../../other')

    def test_asset_revision_and_integrity(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            path = root / 'demo.png'
            path.write_bytes(b'original')
            lib = MediaLibrary(root)
            lib.add(str(path), 'Hand hygiene', 'test', 'owned', 'internal', 'HH', '1', 'reviewer')
            self.assertEqual(1, len(lib.match('Hand hygiene', 'HH', '1')))
            self.assertEqual([], lib.match('Hand hygiene', 'HH', '2'))
            path.write_bytes(b'changed')
            self.assertEqual([], lib.match('Hand hygiene', 'HH', '1'))

    def test_missing_ffmpeg_and_missing_media(self):
        row = storyboard('Wash hands.')[0]
        approve(row, 'QA')
        with patch('sop_studio.composition.ffmpeg', return_value=None):
            with self.assertRaisesRegex(RuntimeError, 'FFmpeg'):
                render([row], 'Test', 'Windows offline', '', 155)
        with self.assertRaisesRegex(ValueError, 'missing'):
            preflight([row], {}, None, None, strict_media=True)
        with self.assertRaisesRegex(ValueError, 'exactly one'):
            preflight([row], {}, None, [])

    def test_cancel_subprocess(self):
        import sys
        with tempfile.TemporaryDirectory() as folder:
            cancel = Path(folder) / 'cancel'
            cancel.touch()
            with self.assertRaises(RenderCancelled):
                run([sys.executable, '-c', 'print(1)'], cancel)

    def test_cancel_running_subprocess(self):
        import sys
        import threading
        import time
        with tempfile.TemporaryDirectory() as folder:
            cancel = Path(folder) / 'cancel'
            timer = threading.Timer(.3, cancel.touch)
            timer.start()
            started = time.monotonic()
            try:
                with self.assertRaises(RenderCancelled):
                    run([sys.executable, '-c', 'import time; time.sleep(30)'], cancel)
            finally:
                timer.cancel()
            self.assertLess(time.monotonic()-started, 3)

    def test_pronunciation_does_not_replace_substrings(self):
        self.assertEqual('G M P GMPx', pronunciation('GMP GMPx', {'GMP': 'G M P'}))

    def test_model_missing_is_actionable(self):
        from sop_studio.llm import enrich_scene
        with patch('sop_studio.llm.models', return_value=[]):
            with self.assertRaisesRegex(ValueError, 'installed'):
                enrich_scene(storyboard('Wash hands.')[0], 'not-installed')

if __name__ == '__main__':
    unittest.main()
